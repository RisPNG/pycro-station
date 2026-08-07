from __future__ import annotations

import argparse
import csv
import re
import threading
from collections import Counter, OrderedDict
from copy import copy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PySide6.QtCore import Qt, Signal
from PySide6.QtWidgets import QFileDialog, QGridLayout, QHBoxLayout, QLabel, QSizePolicy, QTextEdit, QVBoxLayout, QWidget
from qfluentwidgets import MessageBox, PrimaryPushButton


class MainWidget(QWidget):
    log_message = Signal(str)
    processing_done = Signal(bool, str)

    def __init__(self):
        super().__init__()
        self.setObjectName("spa_routine_widget")
        self.selected_paths = {
            "input": "",
            "pco2161": "",
            "pco2195": "",
            "size_list": "",
            "region_destination": "",
            "developers": "",
            "developers_location": "",
        }
        self._build_ui()
        self._connect_signals()
        self._refresh_files_box()

    def _build_ui(self):
        self.desc_label = QLabel("", self)
        self.desc_label.setWordWrap(True)
        self.desc_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        self.desc_label.setAlignment(Qt.AlignLeft | Qt.AlignTop)
        self.desc_label.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.desc_label.setStyleSheet(
            "color: #dcdcdc; background: transparent; padding: 6px; "
            "border: 1px solid #3a3a3a; border-radius: 6px;"
        )
        self.set_long_description("")

        self.select_input_btn = PrimaryPushButton("Select SPA Input", self)
        self.select_pco_btn = PrimaryPushButton("Select PCO2161", self)
        self.select_pco2195_btn = PrimaryPushButton("Select PCO2195", self)
        self.select_size_btn = PrimaryPushButton("Size List (Optional)", self)
        self.select_region_btn = PrimaryPushButton("Region & Destination (Optional)", self)
        self.select_developers_btn = PrimaryPushButton("Developers (Optional)", self)
        self.select_locations_btn = PrimaryPushButton("Developers Location (Optional)", self)
        self.run_btn = PrimaryPushButton("Run", self)
        self.selection_buttons = (
            self.select_input_btn,
            self.select_pco_btn,
            self.select_pco2195_btn,
            self.select_size_btn,
            self.select_region_btn,
            self.select_developers_btn,
            self.select_locations_btn,
        )

        self.files_label = QLabel("Selected files", self)
        self.files_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.files_label.setStyleSheet("color: #dcdcdc; background: transparent; padding-left: 2px;")
        self.logs_label = QLabel("Process logs", self)
        self.logs_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.logs_label.setStyleSheet("color: #dcdcdc; background: transparent; padding-left: 2px;")

        shared_style = (
            "QTextEdit{background: #1f1f1f; color: #d0d0d0; "
            "border: 1px solid #3a3a3a; border-radius: 6px;}"
        )
        self.files_box = QTextEdit(self)
        self.files_box.setReadOnly(True)
        self.files_box.setPlaceholderText("Selected files will appear here")
        self.files_box.setStyleSheet(shared_style)
        self.log_box = QTextEdit(self)
        self.log_box.setReadOnly(True)
        self.log_box.setPlaceholderText("Live process log will appear here")
        self.log_box.setStyleSheet(shared_style)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(16, 16, 16, 16)
        main_layout.setSpacing(12)
        main_layout.addWidget(self.desc_label, 1)

        selectors_layout = QGridLayout()
        selectors_layout.setHorizontalSpacing(12)
        selectors_layout.setVerticalSpacing(12)
        selectors_layout.addWidget(self.select_input_btn, 0, 0)
        selectors_layout.addWidget(self.select_pco_btn, 0, 1)
        selectors_layout.addWidget(self.select_pco2195_btn, 1, 0)
        selectors_layout.addWidget(self.select_size_btn, 1, 1)
        selectors_layout.addWidget(self.select_region_btn, 2, 0)
        selectors_layout.addWidget(self.select_developers_btn, 2, 1)
        selectors_layout.addWidget(self.select_locations_btn, 3, 0)
        selectors_layout.setColumnStretch(0, 1)
        selectors_layout.setColumnStretch(1, 1)
        main_layout.addLayout(selectors_layout)

        run_layout = QHBoxLayout()
        run_layout.addStretch(1)
        run_layout.addWidget(self.run_btn, 1)
        run_layout.addStretch(1)
        main_layout.addLayout(run_layout)

        labels_layout = QHBoxLayout()
        labels_layout.addWidget(self.files_label, 1)
        labels_layout.addWidget(self.logs_label, 1)
        main_layout.addLayout(labels_layout)

        content_layout = QHBoxLayout()
        content_layout.addWidget(self.files_box, 1)
        content_layout.addWidget(self.log_box, 1)
        main_layout.addLayout(content_layout, 4)

    def set_long_description(self, text: str):
        clean = (text or "").strip()
        if clean:
            self.desc_label.setText(clean)
            self.desc_label.show()
        else:
            self.desc_label.clear()
            self.desc_label.hide()

    def _connect_signals(self):
        self.select_input_btn.clicked.connect(lambda: self.select_file("input", "Select SPA Input Workbook"))
        self.select_pco_btn.clicked.connect(lambda: self.select_file("pco2161", "Select PCO2161 Workbook"))
        self.select_pco2195_btn.clicked.connect(lambda: self.select_file("pco2195", "Select PCO2195 Workbook"))
        self.select_size_btn.clicked.connect(lambda: self.select_file("size_list", "Select Size List Workbook"))
        self.select_region_btn.clicked.connect(lambda: self.select_file("region_destination", "Select Region & Destination Workbook"))
        self.select_developers_btn.clicked.connect(lambda: self.select_file("developers", "Select Developers Workbook"))
        self.select_locations_btn.clicked.connect(lambda: self.select_file("developers_location", "Select Developers Location Workbook"))
        self.run_btn.clicked.connect(self.run_process)
        self.log_message.connect(self.append_log)
        self.processing_done.connect(self.on_processing_done)

    def select_file(self, key: str, title: str):
        path, _ = QFileDialog.getOpenFileName(self, title, "", "Spreadsheet Files (*.xlsx *.xlsm *.ods)")
        if path:
            self.selected_paths[key] = path
            self._refresh_files_box()

    def _refresh_files_box(self):
        defaults_dir = Path(__file__).resolve().parent / "defaults"
        labels = (
            ("SPA Input", "input", "Required"),
            ("PCO2161", "pco2161", "Required"),
            ("PCO2195", "pco2195", "Required"),
            ("Size List", "size_list", str(defaults_dir / "Size List.xlsx")),
            ("Region & Destination", "region_destination", str(defaults_dir / "Region & Destination.xlsx")),
            ("Developers", "developers", str(defaults_dir / "Developers.xlsx")),
            ("Developers Location", "developers_location", str(defaults_dir / "Developers Location.xlsx")),
        )
        self.files_box.setPlainText("\n".join(f"{label}: {self.selected_paths[key] or fallback}" for label, key, fallback in labels))

    def run_process(self):
        if not self.selected_paths["input"] or not self.selected_paths["pco2161"] or not self.selected_paths["pco2195"]:
            MessageBox("Warning", "Select the SPA input, PCO2161, and PCO2195 workbooks.", self).exec()
            return
        self.log_box.clear()
        self.log_message.emit("Process starts")
        self.run_btn.setEnabled(False)
        for button in self.selection_buttons:
            button.setEnabled(False)

        def worker():
            try:
                outputs = process_file(
                    self.selected_paths["input"],
                    self.selected_paths["pco2161"],
                    self.selected_paths["pco2195"],
                    self.selected_paths["size_list"] or None,
                    self.selected_paths["region_destination"] or None,
                    self.selected_paths["developers"] or None,
                    self.selected_paths["developers_location"] or None,
                    log_emit=self.log_message.emit,
                )
                self.processing_done.emit(True, "\n".join(outputs))
            except Exception as error:
                self.log_message.emit(f"ERROR: {error}")
                self.processing_done.emit(False, str(error))

        threading.Thread(target=worker, daemon=True).start()

    def append_log(self, text: str):
        self.log_box.append(text)
        self.log_box.ensureCursorVisible()

    def on_processing_done(self, success: bool, message: str):
        self.run_btn.setEnabled(True)
        for button in self.selection_buttons:
            button.setEnabled(True)
        title = "Processing complete" if success else "Processing failed"
        dialog = MessageBox(title, message, self)
        dialog.yesButton.setText("OK")
        dialog.cancelButton.hide()
        dialog.exec()


def get_widget():
    return MainWidget()


HEADERS = (
    "P.O. I.D.", "CUST XREF", "FTY", "CU CD", "STYLE NUMBER", "CLR COD",
    "ORDERQTY", "SIZE", "OGAC", "AF REQ#", "MSC", "FO/ND", "TRACK",
    "Job#SSS", "Destination", "Remarks", "Job#", "New Job#", "Season",
    "Drop Style", "Remark 1", "Remark 2", "Product Type", "Vendor", "Team",
    "Smallest OGAC",
)


def vba_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_import_size(value) -> str:
    size = vba_text(value).replace("-", "")
    match = re.match(r"^(X{2,})(.+)$", size, re.IGNORECASE)
    return f"{len(match.group(1))}X{match.group(2)}" if match else size


def format_output_size(value) -> str:
    size = vba_text(value)
    return f"XX{size[2:]}" if len(size) > 2 and size[:2].upper() == "2X" else size


def format_colourway(value) -> str:
    return re.sub(r"\d+", lambda match: match.group().zfill(3), vba_text(value))


def increment_letter(value: str) -> str:
    result = []
    carry = True
    for character in reversed(value.upper()):
        if carry and character == "Z":
            result.append("A")
        elif carry:
            result.append(chr(ord(character) + 1))
            carry = False
        else:
            result.append(character)
    if carry:
        result.append("A")
    return "".join(reversed(result))


def build_processed_rows(source_ws, size_ws, pco_ws, pco2195_ws, mapping_ws):
    size_values = []
    for row in range(2, size_ws.max_row + 1):
        value = size_ws.cell(row, 1).value
        if value is not None and vba_text(value):
            size_values.append(value)
    if not size_values:
        raise ValueError("The Size List workbook does not contain sizes in column A below the header.")
    size_order = {vba_text(value): index for index, value in enumerate(size_values)}
    raw_rows = []
    for source_row in source_ws.iter_rows(min_row=2, values_only=True):
        row = [None] * 27
        row[1] = source_row[2]
        row[2] = source_row[6]
        row[3] = source_row[2]
        row[4] = source_row[7]
        row[5] = source_row[8]
        row[6] = source_row[9]
        row[7] = source_row[11]
        row[8] = normalize_import_size(source_row[12])
        row[9] = source_row[21]
        row[10] = source_row[67]
        row[12] = source_row[66]
        row[13] = source_row[1]
        row[21] = vba_text(source_row[74])
        row[22] = vba_text(source_row[68])
        raw_rows.append(row)
    raw_rows.sort(
        key=lambda row: (
            (0, float(row[5])) if isinstance(row[5], (int, float)) and not isinstance(row[5], bool) else (1, vba_text(row[5])),
            (0, float(row[6])) if isinstance(row[6], (int, float)) and not isinstance(row[6], bool) else (1, vba_text(row[6])),
            size_order.get(vba_text(row[8]), len(size_order)),
        )
    )

    mapping = {}
    for row in mapping_ws.iter_rows(min_row=2, values_only=True):
        mapping[vba_text(row[0])] = row

    rows = []
    style_groups = OrderedDict()
    for row in raw_rows:
        style_groups.setdefault(vba_text(row[5]), []).append(row)
    for style_rows in style_groups.values():
        colour_groups = OrderedDict()
        for row in style_rows:
            colour_groups.setdefault(vba_text(row[6]), []).append(row)
        for colour_rows in colour_groups.values():
            size_groups = OrderedDict()
            for row in colour_rows:
                size_groups.setdefault(vba_text(row[8]), []).append(row)
            for size_rows in size_groups.values():
                rows.extend(size_rows)
                total_size = [None] * 27
                total_size[0] = "Total Size Qty"
                total_size[5] = size_rows[0][5]
                total_size[6] = size_rows[0][6]
                total_size[7] = sum(float(row[7] or 0) for row in size_rows)
                total_size[8] = size_rows[0][8]
                rows.append(total_size)
            first = colour_rows[0]
            sizes = {vba_text(row[8]) for row in colour_rows}
            fit_type = vba_text(first[22])
            product_type = vba_text(first[21])
            keeping_size = ""
            if fit_type == "WESTERN FIT" and product_type == "MENS":
                keeping_size = "M" if "M" in sizes else "32" if "32" in sizes else ""
            elif fit_type == "WESTERN FIT" and product_type == "WOMENS":
                keeping_size = "S" if "S" in sizes else "6" if "6" in sizes else ""
            elif fit_type == "ASIA FIT" and product_type == "MENS":
                keeping_size = "L" if "L" in sizes else "32" if "32" in sizes else ""
            elif fit_type == "ASIA FIT" and product_type == "WOMENS":
                keeping_size = "M" if "M" in sizes else "6" if "6" in sizes else ""
            elif product_type == "KIDS" and fit_type in {"WESTERN FIT", "ASIA FIT"} and "M" in sizes:
                keeping_size = "M"
            elif fit_type == "PLUS SIZE" and product_type in {"MENS", "WOMENS"} and "2X" in sizes:
                keeping_size = "2X"
            elif fit_type == "PLUS SIZE" and product_type == "KIDS" and "M+" in sizes:
                keeping_size = "M+"
            keeping = [None] * 27
            keeping[0] = "Factory Keeping Qty"
            keeping[5] = first[5]
            keeping[6] = first[6]
            keeping[7] = None if keeping_size else 0
            keeping[8] = keeping_size or None
            keeping[10] = "USAO" if fit_type in {"WESTERN FIT", "PLUS SIZE"} else "APAO"
            keeping[21] = first[21]
            keeping[22] = first[22]
            rows.append(keeping)
            total_colour = [None] * 27
            total_colour[0] = "Total CW Qty"
            total_colour[5] = first[5]
            total_colour[6] = first[6]
            total_colour[7] = sum(float(row[7] or 0) for row in colour_rows)
            total_colour[13] = first[13]
            dates = [row[9] for row in colour_rows if isinstance(row[9], (date, datetime))]
            total_colour[26] = min(dates) if dates else None
            rows.append(total_colour)
        af_req = Counter(vba_text(row[10]) for row in style_rows if vba_text(row[10])).most_common(1)
        style_total = [None] * 27
        style_total[0] = "Total Style Qty"
        style_total[5] = style_rows[0][5]
        if af_req and af_req[0][0] in mapping:
            style_total[15] = mapping[af_req[0][0]][3]
            style_total[16] = mapping[af_req[0][0]][0]
        rows.extend((style_total, [None] * 27))

    pco_rows = [list(row) for row in pco_ws.iter_rows(min_row=6, values_only=True)]
    pco2195_jobs = {
        vba_text(row[0]).upper()
        for row in pco2195_ws.iter_rows(min_col=5, max_col=5, values_only=True)
        if vba_text(row[0]) and vba_text(row[0]).upper() != "JOB NO"
    }
    factory_indexes = [index for index, row in enumerate(rows) if row[0] == "Factory Keeping Qty"]
    for index in factory_indexes:
        row = rows[index]
        if row[7] is None:
            match = next((pco for pco in pco_rows if vba_text(pco[2]) == vba_text(row[5]) and vba_text(pco[10]) == "SIV"), None)
            if match and vba_text(match[4]) in {"A14", "A17"}:
                row[7] = 9
            elif vba_text(row[22]) == "PLUS SIZE":
                row[7] = 1
            elif vba_text(row[21]) in {"KIDS", "MENS", "WOMENS"}:
                row[7] = 6
            else:
                row[7] = 1 if vba_text(row[22]) == "PLUS SIZE" else 6
    for index in factory_indexes:
        row = rows[index]
        if not row[7]:
            continue
        match = next((pco for pco in pco_rows if vba_text(pco[2]) == vba_text(row[5]) and vba_text(pco[10]) == "SIV"), None)
        fit_type = vba_text(row[22])
        remark = vba_text(match[12]) if match else ""
        if not match or fit_type == "PLUS SIZE" or not ((fit_type == "WESTERN FIT" and remark.startswith("AS")) or (fit_type == "ASIA FIT" and remark.startswith("WS"))):
            continue
        linked_style = re.split(r"[ \-/,;=]", remark[2:].lstrip(), maxsplit=1)[0]
        linked_index = next((other for other in factory_indexes if vba_text(rows[other][5]) == linked_style and vba_text(rows[other][6]) == vba_text(row[6])), None)
        if linked_index is None:
            continue
        rows[index + 1][21] = linked_style
        linked_pco = next((pco for pco in pco_rows if vba_text(pco[2]) == linked_style), None)
        if not linked_pco:
            continue
        linked_fit = vba_text(rows[linked_index][22])
        category = vba_text(linked_pco[4])
        if category in {"A14", "A17"}:
            if linked_fit == "WESTERN FIT":
                row[7] = 0
            elif linked_fit == "ASIA FIT":
                row[7] = 9
                rows[linked_index][7] = 0
        elif vba_text(row[21]) in {"KIDS", "MENS", "WOMENS"}:
            if linked_fit == "WESTERN FIT":
                row[7] = 0
            elif linked_fit == "ASIA FIT":
                row[7] = 6
                rows[linked_index][7] = 0

    unique_job = {vba_text(row[5]): False for row in rows}
    for row in rows:
        if row[0] != "Total CW Qty":
            continue
        processed = False
        for pco in pco_rows:
            if vba_text(pco[2]) != vba_text(row[5]) or vba_text(pco[10]) != "SIV":
                continue
            row[17] = vba_text(pco[6]) or None
            row[19] = vba_text(pco[1]) or None
            row[20] = vba_text(pco[11]) or None
            row[25] = vba_text(pco[9]) or None
            row[23] = vba_text(pco[3]) or None
            row[24] = vba_text(pco[10]) or None
            split_job = ""
            colourway = ""
            remark = vba_text(pco[12])
            split_values = [value.strip() for value in remark.split(",")] if remark else []
            combined = "," + ",".join(format_colourway(value.split("=", 1)[1]) for value in split_values if "=" in value)
            skip_pco = False
            for value in split_values:
                if "=" in value and ("#" in value or "OTHER" in value.upper()):
                    split_job, colourway = (part.strip() for part in value.split("=", 1))
                    colourway = format_colourway(colourway)
                current_colour = format_colourway(row[6])
                if current_colour in colourway or (colourway.upper() == "OTHER" and current_colour not in combined):
                    pco_job = vba_text(pco[6])
                    previous_job = vba_text(pco[7])
                    if split_job and ((pco_job and (split_job in pco_job or pco_job in split_job)) or (previous_job and (split_job in previous_job or previous_job in split_job))):
                        unique_job[vba_text(row[5])] = True
                        break
                    if split_job and not re.search(r"[A-Za-z]", colourway):
                        skip_pco = True
                        break
            if skip_pco:
                continue
            if split_job:
                row[17] = split_job
            if not processed:
                processed = True
                pco_job = vba_text(pco[6])
                last_salesman_job = vba_text(pco[8])
                if split_job in pco_job or pco_job in split_job:
                    if not pco_job:
                        row[17] = last_salesman_job.replace("SSS", "")
                        pco_job = vba_text(row[17])
                    if not last_salesman_job:
                        row[14] = re.sub(r"(\d+)", r"\1SSS", pco_job)
                    elif len(pco_job) + 3 == len(last_salesman_job):
                        row[14] = last_salesman_job + "A"
                    else:
                        suffix = last_salesman_job[len(pco_job) + 3:]
                        row[14] = last_salesman_job[:len(pco_job) + 3] + (increment_letter(suffix) if suffix.isalpha() else "A")
                else:
                    if split_job[:2].upper() in {"AS", "WS"}:
                        split_job = split_job[2:].strip()
                    row[14] = re.sub(r"(\d+)", r"\1SSS", split_job)
            if unique_job[vba_text(row[5])]:
                break

    for index in factory_indexes:
        row = rows[index]
        if not row[7] or vba_text(row[22]) == "PLUS SIZE":
            continue
        match = next((pco for pco in pco_rows if vba_text(pco[2]) == vba_text(row[5]) and vba_text(pco[10]) == "SIV"), None)
        if match and vba_text(match[4]) in {"A14", "A17"}:
            continue
        new_job = vba_text(rows[index + 1][14]).upper()
        base_job = vba_text(rows[index + 1][17]).upper()
        if new_job not in pco2195_jobs or not base_job:
            row[7] = 6
            continue
        base_sss_job = re.sub(r"(\d+)", r"\1SSS", base_job)
        if not new_job.startswith(base_sss_job):
            row[7] = 6
            continue
        suffix = new_job[len(base_sss_job):]
        if suffix == "A":
            row[7] = 5
        elif suffix.isalpha() and suffix:
            row[7] = 3
        else:
            row[7] = 6

    last_style_start = 0
    for index, row in enumerate(rows):
        if row[0] == "Total Style Qty":
            jobs = []
            for preceding in reversed(rows[last_style_start:index]):
                if preceding[14] and preceding[14] not in jobs:
                    jobs.append(preceding[14])
            row[14] = " / ".join(jobs) or None
            last_style_start = index + 1
    style_dates = {}
    for index, row in enumerate(rows):
        if row[0] == "Total CW Qty":
            if index and rows[index - 1][0] == "Factory Keeping Qty":
                row[7] = float(row[7] or 0) + float(rows[index - 1][7] or 0)
            if not row[14]:
                row[14] = "BDS NO DATA"
            if isinstance(row[26], (date, datetime)):
                key = vba_text(row[5])
                style_dates[key] = min(style_dates.get(key, row[26]), row[26])
    style_totals = Counter()
    for row in rows:
        if row[0] == "Total CW Qty":
            style_totals[vba_text(row[5])] += row[7]
    for row in rows:
        if row[0] == "Total CW Qty" and vba_text(row[5]) in style_dates:
            row[26] = style_dates[vba_text(row[5])]
        elif row[0] == "Total Style Qty":
            row[7] = style_totals[vba_text(row[5])]
        elif row[0] is None:
            row[21] = None
            row[22] = None
    for index, row in enumerate(rows):
        if row[0] == "Total Style Qty" and not row[14]:
            jobs = []
            scan = index - 1
            while scan >= 0 and rows[scan][0] != "Total Style Qty":
                job = rows[scan][14]
                if job and job not in jobs:
                    jobs.append(job)
                scan -= 1
            row[14] = " / ".join(jobs) or None
    return rows, size_values


def generate_obs_ssppr(rows, size_values, output_path: Path, run_time: datetime):
    prefix = ["Vendor", "Planning Season", "Year", "Material", "Job Number", "Product Type", "VNFOB", "PO Number", "OGAC Date", "Delivery Mode", "Packing Mode", "Job Type", "Status", "ID", "Price Per Unit", "Order Taken Date", "Buy Month", "Team", "Destination", "Track", "Keeping Size"]
    total_indexes = [index for index, row in enumerate(rows) if row[0] == "Total CW Qty"]
    output_rows = []
    for index in total_indexes:
        row = rows[index]
        size_quantities = OrderedDict()
        scan = index - 1
        while scan >= 0 and rows[scan][0] != "Total CW Qty":
            candidate = rows[scan]
            if candidate[0] in {"Total Size Qty", "Factory Keeping Qty"} and vba_text(candidate[5]) == vba_text(row[5]) and vba_text(candidate[6]) == vba_text(row[6]) and candidate[7]:
                size_quantities[vba_text(candidate[8])] = size_quantities.get(vba_text(candidate[8]), 0) + int(candidate[7])
            scan -= 1
        destination = next((vba_text(candidate[15]) for candidate in rows[index:] if vba_text(candidate[15])), "")
        season = vba_text(row[19])
        record = [
            vba_text(row[24]), season[:2], "20" + season[-2:], f"{vba_text(row[5])}-{vba_text(row[6])}",
            vba_text(row[14]), vba_text(row[23]), "N", "TBA",
            row[26].strftime("%m/%d/%Y") if isinstance(row[26], (date, datetime)) else "",
            "NAF", "B", "S", "F", f"{season[:2]}{season[-2:]}{vba_text(row[24])}SS", "10.00",
            run_time.strftime("%m/%d/%Y"), run_time.strftime("%Y%m"), vba_text(row[25]), destination,
            vba_text(row[13]), format_output_size(rows[index - 1][8]) if index else "",
        ]
        total = 0
        for size in size_values:
            quantity = size_quantities.get(vba_text(size), "")
            record.append(quantity)
            if quantity != "":
                total += int(quantity)
        record.extend((total, "-"))
        output_rows.append(record)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(prefix + size_values + ["TOTAL", "-"])
        writer.writerows(output_rows)


def generate_spa_order(rows, mapping_ws, output_path: Path):
    mapping = {}
    for values in mapping_ws.iter_rows(min_row=2, values_only=True):
        key = vba_text(values[0])
        if key and key not in mapping:
            mapping[key] = (vba_text(values[1]), vba_text(values[2]))
    job_lookup = {(vba_text(row[5]), vba_text(row[6])): vba_text(row[14]) for row in rows if row[0] == "Total CW Qty"}
    main = OrderedDict()
    keeping = OrderedDict()
    style_totals = Counter()
    style_jobs = OrderedDict()
    for row in rows:
        style = vba_text(row[5])
        colour = vba_text(row[6])
        if not style:
            continue
        style_jobs.setdefault(style, OrderedDict())
        if row[0] is None and row[1] is not None and isinstance(row[7], (int, float)) and row[7] > 0:
            region, cnc = mapping.get(vba_text(row[10]), (f"RegNotFound_{vba_text(row[10])}", f"CNCNotFound_{vba_text(row[10])}"))
            job = job_lookup.get((style, colour), "JobNoNotFound")
            main.setdefault((style, colour), OrderedDict())
            key = (format_output_size(row[8]), region, cnc, job)
            main[(style, colour)][key] = main[(style, colour)].get(key, 0) + row[7]
            style_totals[style] += row[7]
            style_jobs[style][job] = None
        elif row[0] == "Factory Keeping Qty" and isinstance(row[7], (int, float)) and row[7] > 0:
            region, cnc = ("USAO", "445591") if vba_text(row[22]) in {"WESTERN FIT", "PLUS SIZE"} else ("APAO", "4456076L")
            job = job_lookup.get((style, colour), "JobNoNotFound_K")
            keeping.setdefault((style, colour), []).append((style, colour, region, cnc, job, format_output_size(row[8]), row[7], "Keeping"))
            style_totals[style] += row[7]
            style_jobs[style][job] = None
    group_keys = sorted(set(main) | set(keeping), key=lambda key: f"{key[0]}|{key[1]}")
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\r\n")
        writer.writerow(("Style", "CW", "Region", "CNC IM#", "Job No.", "Size", "Total Order Qty", "Keeping"))
        previous_style = ""
        for style, colour in group_keys:
            if previous_style and previous_style != style:
                writer.writerow((f"{previous_style} Total Qty", "", "", "", " / ".join(style_jobs[previous_style]), "", style_totals[previous_style], ""))
            for (size, region, cnc, job), quantity in main.get((style, colour), {}).items():
                writer.writerow((style, colour, region, cnc, job, size, quantity, ""))
            writer.writerows(keeping.get((style, colour), ()))
            previous_style = style
        if previous_style:
            writer.writerow((f"{previous_style} Total Qty", "", "", "", " / ".join(style_jobs[previous_style]), "", style_totals[previous_style], ""))


def generate_wash_test(rows, source_ws, pco_ws, developers_ws, developers_location_ws, output_path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wash Test Data"
    headers = ("No.", "Category", "Category Desc.", "Job No.", "Style No.", "Style Desc.", "Col. Code", "Keeping Qty", "CW Storage Test Request", "Garment Wash Test Request", "PCC DS Keeping", "Size", "Prod. Loc.", "SS Ex-fty Date", "Wash Sample Rec. Date", "Test Report Ready Date", "Test Report to NSQ", "NSQ Reject", "Resend to NSQ", "Developer", "Cc Email to Developer", "Spec.", "Remarks")
    sheet.append([None] * 23)
    for _ in range(5):
        sheet.append([None] * 23)
    sheet.append(headers)
    pco_rows = list(pco_ws.iter_rows(min_row=6, values_only=True))
    raw_descriptions = {(vba_text(row[8]), vba_text(row[9])): row[15] for row in source_ws.iter_rows(min_row=2, values_only=True)}
    last_style = ""
    number = 0
    developers = OrderedDict()
    for index, row in enumerate(rows):
        if row[0] != "Factory Keeping Qty" or row[7] == 0:
            continue
        style = vba_text(row[5])
        job = vba_text(rows[index + 1][17])
        for pco_index, pco in enumerate(pco_rows):
            previous_job = vba_text(pco_rows[pco_index - 1][6]) if pco_index else ""
            if vba_text(pco[2]) != style or vba_text(pco[6]) != job or vba_text(pco[6]) == previous_job:
                continue
            if style != last_style:
                number += 1
                last_style = style
                display_number = number
            else:
                display_number = None
            developer = vba_text(pco[13])
            developers[developer] = None
            sheet.append((display_number, pco[4], pco[5], rows[index + 1][14], row[5], raw_descriptions.get((style, vba_text(row[6]))), row[6], row[7], None, None, None, format_output_size(row[8]), "VN", rows[index + 2][26], None, None, None, None, None, pco[13], None, row[22], pco[12]))
    sheet.column_dimensions["A"].hidden = True
    sheet.auto_filter.ref = f"A7:W{sheet.max_row}"
    sheet["A7"].fill = PatternFill("solid", fgColor="C8C8FF")
    for cell in sheet[7]:
        cell.fill = PatternFill("solid", fgColor="C8C8FF")
    widths = (6.426, 10.855, 20.57, 14.285, 10.855, 38, 11.57, 13.57, 25, 28.285, 17, 6.711, 12, 14.711, 24.285, 23.426, 19.711, 12.855, 16.285, 16.141, 22.426, 12.285, 68.285)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in range(8, sheet.max_row + 1):
        sheet.cell(row, 14).number_format = "dd/mm/yyyy"
    validation = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"I8:I{sheet.max_row + 1}")
    validation.add(f"J8:J{sheet.max_row + 1}")
    validation.add(f"K8:K{sheet.max_row + 1}")
    group_starts = [row for row in range(8, sheet.max_row + 1) if sheet.cell(row, 1).value is not None]
    medium = Side(style="medium", color="000000")
    hair = Side(style="hair", color="808080")
    for group_index, start in enumerate(group_starts):
        end = group_starts[group_index + 1] - 1 if group_index + 1 < len(group_starts) else sheet.max_row
        for row in range(start, end + 1):
            for column in range(1, 24):
                sheet.cell(row, column).border = Border(
                    left=medium if column == 1 else hair,
                    right=medium if column == 23 else hair,
                    top=medium if row == start else hair,
                    bottom=medium if row == end else hair,
                )

    source_locations = developers_location_ws
    locations = workbook.create_sheet("Developer Locations")
    for row in source_locations.iter_rows():
        for cell in row:
            target = locations[cell.coordinate]
            target.value = cell.value
            if cell.has_style:
                target.font = copy(cell.font)
                target.fill = copy(cell.fill)
                target.border = copy(cell.border)
                target.alignment = copy(cell.alignment)
                target.number_format = cell.number_format
                target.protection = copy(cell.protection)
    for merged in source_locations.merged_cells.ranges:
        locations.merge_cells(str(merged))
    for key, dimension in source_locations.column_dimensions.items():
        locations.column_dimensions[key].width = dimension.width
    locations["A1"] = "Developers"
    locations["B1"] = "Location"
    dev_ws = developers_ws
    output_row = 2
    for developer in developers:
        if not developer:
            continue
        location = ""
        for column in range(1, dev_ws.max_column + 1):
            for row in range(2, dev_ws.max_row + 1):
                if developer in vba_text(dev_ws.cell(row, column).value):
                    location = vba_text(dev_ws.cell(1, column).value)
                    break
            if location:
                break
        locations.cell(output_row, 1, developer)
        locations.cell(output_row, 2, location)
        output_row += 1
    locations.auto_filter.ref = f"A1:B{max(output_row - 1, 1)}"
    for cell in locations[1][:2]:
        cell.fill = PatternFill("solid", fgColor="C8C8FF")
    locations.column_dimensions["A"].width = 16.141
    locations.column_dimensions["B"].width = 12.141
    locations.column_dimensions["D"].width = 12.711
    locations.column_dimensions["E"].width = 35.285
    workbook.save(output_path)


def load_spreadsheet(path: Path, data_only: bool = False):
    if path.suffix.lower() != ".ods":
        return load_workbook(path, data_only=data_only)
    sheets = pd.read_excel(path, sheet_name=None, header=None, engine="odf", keep_default_na=False)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, data in sheets.items():
        sheet = workbook.create_sheet(str(sheet_name)[:31])
        for row_number, row in enumerate(data.itertuples(index=False, name=None), 1):
            for column_number, value in enumerate(row, 1):
                if value == "" or pd.isna(value):
                    continue
                if hasattr(value, "to_pydatetime"):
                    value = value.to_pydatetime()
                elif hasattr(value, "item"):
                    value = value.item()
                sheet.cell(row_number, column_number, value)
    if not workbook.sheetnames:
        workbook.create_sheet("Sheet1")
    return workbook


def process_file(
    input_path: str,
    pco2161_path: str,
    pco2195_path: str,
    size_list_path: str | None = None,
    region_destination_path: str | None = None,
    developers_path: str | None = None,
    developers_location_path: str | None = None,
    output_dir: str | None = None,
    run_time: datetime | None = None,
    log_emit=None,
):
    source_path = Path(input_path)
    if not source_path.is_file():
        raise FileNotFoundError(f"File not found: {source_path}")
    pco_path = Path(pco2161_path)
    if not pco_path.is_file():
        raise FileNotFoundError(f"PCO2161 file not found: {pco_path}")
    pco2195_file_path = Path(pco2195_path)
    if not pco2195_file_path.is_file():
        raise FileNotFoundError(f"PCO2195 file not found: {pco2195_file_path}")
    run_time = run_time or datetime.now()
    destination = Path(output_dir) if output_dir else source_path.parent
    destination.mkdir(parents=True, exist_ok=True)
    defaults_dir = Path(__file__).resolve().parent / "defaults"
    size_path = Path(size_list_path) if size_list_path else defaults_dir / "Size List.xlsx"
    mapping_path = Path(region_destination_path) if region_destination_path else defaults_dir / "Region & Destination.xlsx"
    developers_file_path = Path(developers_path) if developers_path else defaults_dir / "Developers.xlsx"
    locations_path = Path(developers_location_path) if developers_location_path else defaults_dir / "Developers Location.xlsx"
    for label, path in (
        ("Size List", size_path),
        ("Region & Destination", mapping_path),
        ("Developers", developers_file_path),
        ("Developers Location", locations_path),
    ):
        if not path.is_file():
            raise FileNotFoundError(f"{label} file not found: {path}")
    if callable(log_emit):
        log_emit(f"Reading {source_path.name}")
        log_emit(f"PCO2161: {pco_path.name}")
        log_emit(f"PCO2195: {pco2195_file_path.name}")
        log_emit(f"Size List: {size_path.name}{' (default)' if size_list_path is None else ''}")
        log_emit(f"Region & Destination: {mapping_path.name}{' (default)' if region_destination_path is None else ''}")
        log_emit(f"Developers: {developers_file_path.name}{' (default)' if developers_path is None else ''}")
        log_emit(f"Developers Location: {locations_path.name}{' (default)' if developers_location_path is None else ''}")
    workbook = load_spreadsheet(source_path)
    if "PROCESSED" in workbook.sheetnames:
        raise ValueError("This workbook already contains a PROCESSED sheet.")
    pco_wb = load_spreadsheet(pco_path, data_only=True)
    pco2195_wb = load_spreadsheet(pco2195_file_path, data_only=True)
    size_wb = load_spreadsheet(size_path, data_only=True)
    mapping_wb = load_spreadsheet(mapping_path, data_only=True)
    developers_wb = load_spreadsheet(developers_file_path, data_only=True)
    locations_wb = load_spreadsheet(locations_path, data_only=True)
    pco_ws = pco_wb["pco2161"] if "pco2161" in pco_wb.sheetnames else pco_wb.active
    pco2195_ws = pco2195_wb["pco2195"] if "pco2195" in pco2195_wb.sheetnames else pco2195_wb.active
    size_ws = size_wb["Size List"] if "Size List" in size_wb.sheetnames else size_wb.active
    mapping_ws = mapping_wb["Region & Destination"] if "Region & Destination" in mapping_wb.sheetnames else mapping_wb.active
    developers_ws = developers_wb["Developers"] if "Developers" in developers_wb.sheetnames else developers_wb.active
    locations_ws = locations_wb["Developers Location"] if "Developers Location" in locations_wb.sheetnames else locations_wb.active
    rows, size_values = build_processed_rows(workbook.worksheets[0], size_ws, pco_ws, pco2195_ws, mapping_ws)
    processed = workbook.create_sheet("PROCESSED", 1)
    for column, header in enumerate(HEADERS, 2):
        processed.cell(2, column, header)
    for row_number, values in enumerate(rows, 3):
        for column, value in enumerate(values, 1):
            processed.cell(row_number, column, format_output_size(value) if column == 9 and value is not None else value)
    last_row = len(rows) + 2
    processed.auto_filter.ref = f"A2:AA{last_row - 1}"
    processed.freeze_panes = "A3"
    thin = Side(style="thin", color="000000")
    fills = {"Total Size Qty": "FFFFCC", "Factory Keeping Qty": "ADD8E6", "Total CW Qty": "FFCCCC", "Total Style Qty": "CCFFCC"}
    for row in processed.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=27):
        label = row[0].value
        fill = "C8C8FF" if row[0].row == 2 else fills.get(label, "F0F0F0")
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.fill = PatternFill("solid", fgColor=fill)
        for cell in row[:10]:
            if row[0].row >= 3:
                cell.font = Font(name=cell.font.name or "Calibri", size=cell.font.sz or 11, bold=True)
        if row[0].row >= 3:
            row[14].font = Font(name=row[14].font.name or "Calibri", size=16, bold=True)
        for cell in list(row[1:14]) + list(row[15:21]) + list(row[23:27]):
            cell.alignment = Alignment(horizontal="center")
        row[9].number_format = "mm-dd-yy"
        row[26].number_format = "mm/dd/yyyy"
    for row_number in range(3, last_row):
        processed.row_dimensions[row_number].height = 21
    widths = {1: 19.14, 2: 10.43, 3: 12.43, 4: 6.29, 5: 8.57, 6: 16.29, 7: 10.71, 8: 12.71, 9: 6.86, 10: 10.43, 12: 7, 14: 8.86, 15: 37.43, 16: 13.43, 17: 11, 18: 10, 19: 11.57, 20: 9.71, 21: 12, 22: 11.43, 23: 12.29, 24: 14.57, 25: 9.57, 26: 8.14, 27: 16.43}
    for column, width in widths.items():
        processed.column_dimensions[get_column_letter(column)].width = width
    stamp = run_time.strftime("%Y%m%d_%H%M%S")
    obs_path = destination / f"OBS_SSPPR_{stamp}.csv"
    spa_path = destination / f"SPA_ORDER_{stamp}.csv"
    wash_path = destination / f"WASH_TEST_{stamp}.xlsx"
    processed_path = destination / f"{source_path.stem}_PROCESSED_{stamp}.xlsx"
    generate_obs_ssppr(rows, size_values, obs_path, run_time)
    generate_spa_order(rows, mapping_ws, spa_path)
    generate_wash_test(rows, workbook.worksheets[0], pco_ws, developers_ws, locations_ws, wash_path)
    workbook.save(processed_path)
    pco_wb.close()
    pco2195_wb.close()
    size_wb.close()
    mapping_wb.close()
    developers_wb.close()
    locations_wb.close()
    if callable(log_emit):
        for path in (obs_path, spa_path, wash_path, processed_path):
            log_emit(f"Created {path.name}")
    return tuple(str(path) for path in (obs_path, spa_path, wash_path, processed_path))


def main():
    parser = argparse.ArgumentParser(description="Run the SPA Routine v3.4.0 workflow.")
    parser.add_argument("input", help="Input Excel workbook")
    parser.add_argument("--pco2161", required=True, help="PCO2161 workbook")
    parser.add_argument("--pco2195", required=True, help="PCO2195 workbook")
    parser.add_argument("--size-list", help="Optional Size List workbook")
    parser.add_argument("--region-destination", help="Optional Region & Destination workbook")
    parser.add_argument("--developers", help="Optional Developers workbook")
    parser.add_argument("--developers-location", help="Optional Developers Location workbook")
    parser.add_argument("--output-dir", help="Folder for generated files")
    arguments = parser.parse_args()
    for output in process_file(
        arguments.input,
        arguments.pco2161,
        arguments.pco2195,
        arguments.size_list,
        arguments.region_destination,
        arguments.developers,
        arguments.developers_location,
        arguments.output_dir,
    ):
        print(output)


if __name__ == "__main__":
    main()
