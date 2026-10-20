from __future__ import annotations

import importlib.util
import re
from collections import Counter, OrderedDict, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MONTH_PATTERN = re.compile(r"^\s*(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)'(\d{2})\s*$", re.IGNORECASE)
MONTH_NUMBER = {month: index for index, month in enumerate(
    ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), start=1,
)}
QUANTITY_TOLERANCE = 0.000001
AMOUNT_TOLERANCE = 0.005
DELAY_REASON_LABELS = {
    21: "Late change in tooling spec, trim, design, material",
    83: "Material delay / VCR related issue",
    16: "Divert PO/Consolidate shipments/End-Quarter Inventory Management by Nike Geo",
    3: "Logistic issue from factory",
    "#": "Proposed early shipment rejected by Nike",
    "PG": "Patagonia",
    84: "Quality issue from factory",
    86: "Nike Over Forecast",
    31: "Vessel Delay",
    45: "Production issue from factory",
}


def _read_variance_layout(ws, months: list[str]) -> dict:
    sheet_rows = list(ws.iter_rows(
        min_row=1, max_row=min(ws.max_row, 200), max_col=min(ws.max_column, 40), values_only=True,
    ))
    month_rows = {}
    analysis_header = None
    detail_header = None
    for row_number, row in enumerate(sheet_rows, start=1):
        month = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if month in months:
            month_rows[month] = row_number
        if len(row) < 8 or str(row[6]).strip().casefold() != "due to":
            continue
        if not any(str(value).strip() in months for value in row[7:] if value is not None):
            continue
        if month.casefold() == "month":
            analysis_header = row_number
        else:
            detail_header = row_number

    if len(month_rows) != len(months) or analysis_header is None or detail_header is None:
        raise ValueError("The variance sheet does not contain the selected month rows and analysis tables.")

    analysis_columns = {
        str(value).strip(): column
        for column, value in enumerate(sheet_rows[analysis_header - 1], start=1)
        if value is not None and str(value).strip() in months
    }
    detail_columns = {
        str(value).strip(): column
        for column, value in enumerate(sheet_rows[detail_header - 1], start=1)
        if value is not None and str(value).strip() in months
    }
    if len(analysis_columns) != len(months) or len(detail_columns) != len(months):
        raise ValueError("The variance analysis tables are missing a selected month column.")

    analysis_rows = {}
    for row_number in range(analysis_header + 1, len(sheet_rows) + 1):
        label = sheet_rows[row_number - 1][6]
        if label is None:
            continue
        label = str(label).strip()
        if label.casefold() == "total":
            break
        analysis_rows[label] = row_number
    price_rows = [row for label, row in analysis_rows.items() if label.casefold().startswith("price")]
    delay_rows = [row for label, row in analysis_rows.items() if label.casefold().startswith("delay")]
    early_rows = [row for label, row in analysis_rows.items() if label.casefold().startswith("early")]
    other_rows = [row for label, row in analysis_rows.items()
                  if not label.casefold().startswith(("price", "delay", "early"))]
    if any(len(rows) != 1 for rows in (price_rows, delay_rows, early_rows, other_rows)):
        raise ValueError("The variance analysis needs one price, other, delay, and early row.")

    detail_rows = {}
    for row_number in range(detail_header + 1, len(sheet_rows) + 1):
        label = sheet_rows[row_number - 1][6]
        if label is None:
            continue
        label = str(label).strip()
        if label.casefold() == "total":
            break
        detail_rows[label.casefold()] = row_number
    missing_reasons = [label for label in DELAY_REASON_LABELS.values() if label.casefold() not in detail_rows]
    if missing_reasons:
        raise ValueError("The variance sheet is missing delay reasons: " + ", ".join(missing_reasons))

    return {
        "month_rows": month_rows,
        "analysis_columns": analysis_columns,
        "detail_columns": detail_columns,
        "analysis_rows": {
            "price": price_rows[0], "other": other_rows[0],
            "delay": delay_rows[0], "early": early_rows[0],
        },
        "detail_rows": {code: detail_rows[label.casefold()] for code, label in DELAY_REASON_LABELS.items()},
    }


def _read_sales_comparisons(path: Path, log: Callable[[str], None]) -> tuple[OrderedDict, dict]:
    log(f"Reading Sales Report: {path.name}")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        month_sheets = []
        for ws in wb.worksheets:
            match = MONTH_PATTERN.fullmatch(ws.title)
            if match:
                month = f"{match.group(1).title()}'{match.group(2)}"
                month_sheets.append((2000 + int(match.group(2)), MONTH_NUMBER[match.group(1).title()], month, ws))
        if not month_sheets:
            raise ValueError("Sales Report has no month sheets such as Oct'26.")
        month_sheets.sort()

        comparisons = OrderedDict()
        for _year, _month_number, month, ws in month_sheets:
            heading = next(ws.iter_rows(min_row=1, max_row=1, max_col=4, values_only=True))
            if tuple(str(value).strip().casefold() for value in heading) != ("job", "qty", "amt", "remark"):
                raise ValueError(f"{ws.title} must have Job, Qty, Amt, Remark in columns A:D.")
            jobs = {}
            for row_number, (raw_job, raw_qty, raw_amount, raw_side) in enumerate(
                ws.iter_rows(min_row=2, max_col=4, values_only=True), start=2,
            ):
                if raw_job is None and raw_side is None:
                    continue
                side = str(raw_side).strip().title() if raw_side is not None else ""
                if side not in ("New", "Old"):
                    raise ValueError(f"{ws.title} row {row_number} has an unexpected Old/New marker: {raw_side!r}.")
                job = str(raw_job).strip().upper() if raw_job is not None else ""
                if not job:
                    raise ValueError(f"{ws.title} row {row_number} has no job number.")
                values = jobs.setdefault(job, {}).setdefault(side, [0.0, 0.0])
                values[0] += float(raw_qty) if raw_qty is not None else 0.0
                values[1] += float(raw_amount) if raw_amount is not None else 0.0
            comparisons[month] = jobs
            log(f"{month}: {len(jobs):,} jobs")

        variance_sheets = [ws for ws in wb.worksheets if re.search(r"\bVariance FY\d+\b", ws.title, re.IGNORECASE)]
        if len(variance_sheets) != 1:
            raise ValueError("Sales Report must have one Variance FY sheet.")
        layout = _read_variance_layout(variance_sheets[0], list(comparisons))

        if "Reason code" not in wb:
            raise ValueError("Sales Report has no Reason code sheet.")
        return comparisons, layout
    finally:
        wb.close()


def _read_fotp_gac_codes(path: Path, log: Callable[[str], None]) -> tuple[dict[str, float], dict[str, set]]:
    log(f"Reading FOTP by Season: {path.name}")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "Detail" not in wb:
            raise ValueError("FOTP by Season has no Detail sheet.")
        ws = wb["Detail"]
        header_row = None
        for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=13, values_only=True), start=1):
            if (str(row[1]).strip().casefold() == "siv job no."
                    and str(row[10]).strip().casefold() == "po qty"
                    and str(row[12]).strip().casefold() == "gac code"):
                header_row = row_number
                break
        if header_row is None:
            raise ValueError("FOTP Detail must have SIV Job No. in B, PO Qty in K, and GAC Code in M.")

        quantities = defaultdict(float)
        codes = defaultdict(set)
        for row in ws.iter_rows(min_row=header_row + 1, max_col=13, values_only=True):
            if row[1] is None or row[10] is None:
                continue
            job = str(row[1]).split("-", 1)[0].strip().upper()
            if not job:
                continue
            quantities[job] += float(row[10])
            raw_code = row[12]
            if raw_code is None or not str(raw_code).strip():
                continue
            if isinstance(raw_code, (int, float)) and float(raw_code).is_integer():
                code = int(raw_code)
            elif isinstance(raw_code, str) and raw_code.strip().isdigit():
                code = int(raw_code.strip())
            else:
                code = str(raw_code).strip().upper()
            codes[job].add(code)
        log(f"FOTP: {len(quantities):,} base jobs")
        return quantities, codes
    finally:
        wb.close()


def _read_shipment_statuses(path: Path, log: Callable[[str], None]) -> dict[str, dict[str, float]]:
    log(f"Reading Shipment Forecast: {path.name}")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "SHIPMENTS" not in wb:
            raise ValueError("Shipment Forecast has no SHIPMENTS sheet.")
        ws = wb["SHIPMENTS"]
        header_row = None
        for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=21, values_only=True), start=1):
            if (str(row[4]).strip().casefold() == "job no."
                    and str(row[5]).strip().casefold() == "qty"
                    and str(row[20]).strip().casefold() == "plan ex-fty"):
                header_row = row_number
                break
        if header_row is None:
            raise ValueError("Shipment Forecast must have Job No. in E, Qty in F, and PLAN EX-FTY in U.")

        statuses = defaultdict(lambda: defaultdict(float))
        for row in ws.iter_rows(min_row=header_row + 1, max_col=21, values_only=True):
            status = str(row[20]).strip().upper() if row[20] is not None else ""
            if status not in ("SHORTSHIP", "TBA") or row[4] is None or row[5] is None:
                continue
            job = str(row[4]).split("-", 1)[0].strip().upper()
            statuses[job][status] += float(row[5])
        log(f"Shipment Forecast: {len(statuses):,} shortship/TBA jobs")
        return statuses
    finally:
        wb.close()


def _read_previous_bds_jobs(path: Path, log: Callable[[str], None]) -> set[str]:
    log(f"Reading previous-week BDS: {path.name}")
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "pcp2012" not in wb:
            raise ValueError("Previous-week BDS has no pcp2012 sheet.")
        ws = wb["pcp2012"]
        header_row = None
        job_column = None
        for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=15, values_only=True), start=1):
            matches = [column for column, value in enumerate(row, start=1)
                       if str(value).strip().casefold() == "job number"]
            if matches:
                header_row = row_number
                job_column = matches[-1]
                break
        if header_row is None or job_column is None:
            raise ValueError("Previous-week BDS has no Job Number header in pcp2012.")

        jobs = set()
        for row in ws.iter_rows(min_row=header_row + 1, min_col=job_column, max_col=job_column, values_only=True):
            if row[0] is not None and str(row[0]).strip():
                jobs.add(str(row[0]).strip().upper())
        log(f"Previous-week BDS: {len(jobs):,} jobs")
        return jobs
    finally:
        wb.close()


def _build_recon_sheets(
    comparisons: OrderedDict,
    fotp_quantities: dict[str, float],
    fotp_codes: dict[str, set],
    shipment_statuses: dict[str, dict[str, float]],
    previous_bds_jobs: set[str],
    log: Callable[[str], None],
) -> OrderedDict:
    recon_sheets = OrderedDict()
    by_job = defaultdict(list)
    unresolved = {}
    for month, jobs in comparisons.items():
        rows = []
        for job in sorted(jobs):
            new = jobs[job].get("New")
            old = jobs[job].get("Old")
            row = {
                "job": job,
                "new_qty": new[0] if new is not None else None,
                "new_amount": new[1] if new is not None else None,
                "old_qty": old[0] if old is not None else None,
                "old_amount": old[1] if old is not None else None,
                "remark": None,
                "gac": None,
            }
            qty_difference = (row["new_qty"] or 0.0) - (row["old_qty"] or 0.0)
            amount_difference = (row["new_amount"] or 0.0) - (row["old_amount"] or 0.0)
            if abs(qty_difference) <= QUANTITY_TOLERANCE and abs(amount_difference) <= AMOUNT_TOLERANCE:
                pass
            elif abs(qty_difference) <= QUANTITY_TOLERANCE:
                row["remark"] = "Price Discrepancy"
            elif "SS" in job or re.search(r"\bSAMPLES?\b", job):
                row["remark"] = "Salesman Sample"
            elif len(job) > 5 and job[5].isalpha():
                row["remark"] = "Demand Pull"
            elif job[5:8] == "000":
                row["remark"] = "Forecast PO"
            else:
                unresolved[(month, job)] = row
            rows.append(row)
            by_job[job].append((month, row))
        recon_sheets[f"Recon {month}"] = rows

    candidates = defaultdict(list)
    for job, entries in by_job.items():
        pending = [(month, row) for month, row in entries if (month, job) in unresolved]
        for left_index, (left_month, left_row) in enumerate(pending):
            left_qty = (left_row["new_qty"] or 0.0) - (left_row["old_qty"] or 0.0)
            for right_month, right_row in pending[left_index + 1:]:
                right_qty = (right_row["new_qty"] or 0.0) - (right_row["old_qty"] or 0.0)
                if abs(left_qty) <= QUANTITY_TOLERANCE or abs(left_qty + right_qty) > QUANTITY_TOLERANCE:
                    continue
                candidates[(left_month, job)].append((right_month, job))
                candidates[(right_month, job)].append((left_month, job))

    ambiguous = [(month, job) for (month, job), matches in candidates.items() if len(matches) > 1]
    if ambiguous:
        examples = ", ".join(f"{job} in {month}" for month, job in ambiguous[:5])
        raise ValueError(f"Multiple month movements match the same job and quantity: {examples}.")

    month_positions = {month: position for position, month in enumerate(comparisons)}
    for (month, job), matches in candidates.items():
        other_month, _other_job = matches[0]
        if month_positions[month] > month_positions[other_month]:
            continue
        early_month = month
        later_month = other_month
        early_row = unresolved[(early_month, job)]
        early_qty = (early_row["new_qty"] or 0.0) - (early_row["old_qty"] or 0.0)
        if early_qty > 0:
            remark = f"Early ship fr {later_month} to {early_month}"
        else:
            remark = f"Delay ship fr {early_month} to {later_month}"
        early_row["remark"] = remark
        unresolved[(later_month, job)]["remark"] = remark

    ambiguous_statuses = []
    for (month, job), row in unresolved.items():
        if row["remark"] is not None:
            continue
        qty_difference = (row["new_qty"] or 0.0) - (row["old_qty"] or 0.0)
        matches = [status for status, quantity in shipment_statuses.get(job, {}).items()
                   if abs(abs(qty_difference) - abs(quantity)) <= QUANTITY_TOLERANCE]
        if len(matches) == 1:
            row["remark"] = "shortship" if matches[0] == "SHORTSHIP" else "TBA"
        elif len(matches) > 1:
            ambiguous_statuses.append(f"{job} in {month}")
        if row["remark"] is None:
            if job not in previous_bds_jobs:
                row["remark"] = "Forecast PO New"
            else:
                row["remark"] = "Missing"
    if ambiguous_statuses:
        log(f"Shortship/TBA status is ambiguous for {len(ambiguous_statuses)} jobs: "
            + ", ".join(ambiguous_statuses[:5]))

    unmatched_fotp = 0
    conflicting_gac = 0
    for rows in recon_sheets.values():
        for row in rows:
            if not str(row["remark"] or "").startswith("Delay ship fr "):
                continue
            job = row["job"]
            qty_difference = (row["new_qty"] or 0.0) - (row["old_qty"] or 0.0)
            if abs(fotp_quantities.get(job, 0.0) - abs(qty_difference)) > QUANTITY_TOLERANCE:
                unmatched_fotp += 1
                continue
            codes = fotp_codes.get(job, set())
            if len(codes) == 1:
                row["gac"] = next(iter(codes))
            elif len(codes) > 1:
                conflicting_gac += 1
    if unmatched_fotp:
        log(f"FOTP quantity did not match {unmatched_fotp:,} delay rows; GAC codes remain blank")
    if conflicting_gac:
        log(f"FOTP has differing GAC codes for {conflicting_gac:,} delay rows; GAC codes remain blank")
    counts = Counter(row["remark"] or "Unchanged" for rows in recon_sheets.values() for row in rows)
    log(f"Reconciled {sum(counts.values()):,} month/job rows across {len(recon_sheets)} months")
    return recon_sheets


def _build_variance_updates(
    recon_sheets: OrderedDict,
    layout: dict,
    log: Callable[[str], None],
) -> dict[str, float]:
    updates = {}
    unknown_reasons = set()
    for sheet_name, rows in recon_sheets.items():
        month = sheet_name.removeprefix("Recon ")
        total_row = layout["month_rows"][month]
        updates[f"C{total_row}"] = sum(row["new_amount"] or 0.0 for row in rows)
        updates[f"D{total_row}"] = sum(row["old_amount"] or 0.0 for row in rows)

        category_amounts = {category: 0.0 for category in layout["analysis_rows"]}
        detail_amounts = {code: 0.0 for code in layout["detail_rows"]}
        for row in rows:
            difference = (row["new_amount"] or 0.0) - (row["old_amount"] or 0.0)
            remark = row["remark"] or ""
            if remark == "Price Discrepancy":
                category_amounts["price"] += difference
            elif remark.startswith("Delay ship fr "):
                category_amounts["delay"] += difference
                code = row["gac"]
                if code in detail_amounts:
                    detail_amounts[code] += difference
                elif code is not None:
                    unknown_reasons.add(code)
            elif remark.startswith("Early ship fr "):
                category_amounts["early"] += difference
            else:
                category_amounts["other"] += difference

        analysis_column = get_column_letter(layout["analysis_columns"][month])
        for category, row_number in layout["analysis_rows"].items():
            updates[f"{analysis_column}{row_number}"] = category_amounts[category]
        detail_column = get_column_letter(layout["detail_columns"][month])
        for code, row_number in layout["detail_rows"].items():
            updates[f"{detail_column}{row_number}"] = detail_amounts[code]

    if unknown_reasons:
        log("GAC codes without a matching variance delay reason: "
            + ", ".join(str(code) for code in sorted(unknown_reasons, key=str)))
    return updates


def process_reconciliation(
    sales_path: str | Path,
    fotp_path: str | Path,
    shipment_path: str | Path,
    bds_path: str | Path,
    output_dir: str | Path,
    log: Callable[[str], None],
) -> str:
    paths = [Path(path) for path in (sales_path, fotp_path, shipment_path, bds_path)]
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError(f"Expected an .xlsx workbook: {path}")

    comparisons, layout = _read_sales_comparisons(paths[0], log)
    shipment_statuses = _read_shipment_statuses(paths[2], log)
    previous_bds_jobs = _read_previous_bds_jobs(paths[3], log)
    fotp_quantities, fotp_codes = _read_fotp_gac_codes(paths[1], log)
    recon_sheets = _build_recon_sheets(
        comparisons, fotp_quantities, fotp_codes, shipment_statuses, previous_bds_jobs, log,
    )
    variance_updates = _build_variance_updates(recon_sheets, layout, log)

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now().replace(microsecond=0)
    while True:
        output_path = destination / f"Sales_Forecast_Reconcile_{generated_at:%Y%m%d_%H%M%S}.xlsx"
        if not output_path.exists():
            break
        generated_at += timedelta(seconds=1)

    writer_spec = importlib.util.spec_from_file_location(
        "sales_forecast_reconcile_xlsx_package", Path(__file__).with_name("xlsx_package.py"),
    )
    if writer_spec is None or writer_spec.loader is None:
        raise ImportError("Could not load the Sales Forecast reconciliation workbook writer")
    writer = importlib.util.module_from_spec(writer_spec)
    writer_spec.loader.exec_module(writer)
    writer.write_reconciliation_workbook(
        paths[0], output_path, recon_sheets, variance_updates, ("New", "Old"),
    )
    log(f"Saved Sales Forecast Reconcile: {output_path}")
    return str(output_path)
