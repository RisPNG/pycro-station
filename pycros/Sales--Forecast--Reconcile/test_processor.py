import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


spec = importlib.util.spec_from_file_location(
    "sales_forecast_processor", Path(__file__).with_name("processor.py")
)
processor = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = processor
spec.loader.exec_module(processor)


class SalesForecastReconciliationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.directory.cleanup)
        folder = Path(cls.directory.name)
        cls.sales_path = folder / "Sales Report.xlsx"
        cls.fotp_path = folder / "FOTP by season.xlsx"
        cls.shipment_path = folder / "Shipment forecast.xlsx"
        cls.bds_path = folder / "bds last week.xlsx"

        sales = Workbook()
        variance = sales.active
        variance.title = "45b.Variance FY27"
        variance["B29"] = "May'26"
        variance["B25"] = "FY27"
        variance["B27"] = "Month"
        variance["C27"] = "This Week"
        variance["D27"] = "Last Week"
        variance["G27"] = "Due to"
        variance["G32"] = "Total"
        variance["G46"] = "Due to"
        variance["G57"] = "Total"
        variance["C29"] = 1234
        variance["D29"] = 1200
        variance["B35"] = "Oct'26"
        variance["C35"] = 999
        variance["D35"] = 888
        variance["E35"] = "=C35-D35"
        variance["B37"] = "Nov'26"
        variance["C37"] = 999
        variance["D37"] = 888
        variance["E37"] = "=C37-D37"
        for row, label in {
            28: "Price Revision",
            29: "TBA/Demand Pull/Fx/Forecast PO/DN&CN",
            30: "Delay shipment",
            31: "Early shipment",
            47: "Late change in tooling spec, trim, design, material",
            48: "Material delay / VCR related issue",
            49: "Divert PO/Consolidate shipments/End-Quarter Inventory Management by Nike Geo",
            50: "Logistic issue from factory",
            51: "Proposed early shipment rejected by Nike",
            52: "Patagonia",
            53: "Quality issue from factory",
            54: "Nike Over Forecast",
            55: "Vessel Delay",
            56: "Production issue from factory",
        }.items():
            variance.cell(row, 7, label)
        for row in (27, 46):
            variance.cell(row, 8, "Oct'26")
            variance.cell(row, 9, "Nov'26")
            variance.cell(row, 10, "Dec'26")
        variance["J28"] = 777

        reason = sales.create_sheet("Reason code")
        reason.append([])
        reason.append(["Code #", "Code Description", "SIG Finance Comment"])
        reason.append([21, "N-Dev Delay", "Late change in tooling spec, trim, design, material"])
        reason.append([16, "N-GeoRqHold/Div/Adj", "Divert PO/Consolidate shipments/End-Quarter Inventory Management by Nike Geo"])
        reason.append([3, "N-Logistics", "Logistic issue from Nike"])
        reason.append(["#", "", "Proposed early shipment rejected by Nike"])

        october = sales.create_sheet("Oct'26")
        october.append(["Job", "Qty", "Amt", "Remark"])
        for row in (
            ("AA888888MJ", 10, 100, "Old"),
            ("AA222222MJ", 11, 110, "Old"),
            ("AA333333MJ", 6, 60, "Old"),
            ("AA444444MJ", 8, 80, "Old"),
            ("AA555555MJ", 10, 100, "Old"),
            ("AA666666MJ", 4, 40, "Old"),
            ("AA737373MJ", 2, 20, "Old"),
            ("AA979797MJ", 3, 30, "Old"),
            ("AC123456MJ", 5, 50, "Old"),
            ("AC234567MJ", 8, 80, "Old"),
            ("AA888888MJ", 10, 120, "New"),
            ("AJ005SSSLS", 5, 75, "New"),
            ("BA125MJJA26A3", 8, 80, "New"),
            ("AJ089000MP", 12, 120, "New"),
            ("AJ005010MJ", 9, 90, "New"),
            ("AA111111MJ", 7, 70, "New"),
            ("AC123456MJ", 7, 70, "New"),
        ):
            october.append(row)

        november = sales.create_sheet("Nov'26")
        november.append(["Job", "Qty", "Amt", "Remark"])
        for row in (
            ("AA111111MJ", 7, 70, "Old"),
            ("AJ005SSSLS", 5, 75, "Old"),
            ("BA125MJJA26A3", 8, 80, "Old"),
            ("AJ089000MP", 12, 120, "Old"),
            ("AA222222MJ", 11, 110, "New"),
            ("AA555555MJ", 10, 100, "New"),
            ("AA666666MJ", 4, 40, "New"),
            ("AA737373MJ", 2, 20, "New"),
            ("AA979797MJ", 3, 30, "New"),
        ):
            november.append(row)
        sales.create_sheet("Archive")["A1"] = "leave this sheet alone"
        sales.save(cls.sales_path)

        fotp = Workbook()
        detail = fotp.active
        detail.title = "Detail"
        detail.cell(3, 2, "SIV Job No.")
        detail.cell(3, 11, "PO Qty")
        detail.cell(3, 13, "GAC Code")
        for row_number, job, qty, code in (
            (4, "AA222222MJ-1", 5, 21),
            (5, "AA222222MJ-2", 3, 21),
            (6, "AA222222MJ-3", 3, 21),
            (7, "AA555555MJ-1", 4, 21),
            (8, "AA555555MJ-2", 6, 16),
            (9, "AA666666MJ-1", 2, "#"),
            (10, "AA666666MJ-2", 2, "#"),
            (11, "AA737373MJ", 2, 3),
            (12, "AA979797MJ", 3, 31),
        ):
            detail.cell(row_number, 2, job)
            detail.cell(row_number, 11, qty)
            detail.cell(row_number, 13, code)
        fotp.save(cls.fotp_path)

        forecast = Workbook()
        shipments = forecast.active
        shipments.title = "SHIPMENTS"
        shipments["A1"] = "EXPORT SHIPMENT FORECAST"
        shipments.cell(5, 4, "JOB NO.")
        shipments.cell(5, 5, "JOB NO.")
        shipments.cell(5, 6, "QTY")
        shipments.cell(5, 21, "PLAN EX-FTY")
        shipments.cell(6, 4, "AA333333MJ-short")
        shipments.cell(6, 5, "AA333333MJ")
        shipments.cell(6, 6, 6)
        shipments.cell(6, 21, "shortship")
        shipments.cell(7, 4, "AA444444MJ")
        shipments.cell(7, 5, "AA444444MJ")
        shipments.cell(7, 6, 8)
        shipments.cell(7, 21, "TBA")
        forecast.save(cls.shipment_path)

        bds = Workbook()
        orders = bds.active
        orders.title = "pcp2012"
        orders["A1"] = "Date Prepared : 06/10/2026"
        orders["A2"] = "Order Control List"
        orders["A4"] = "Print by GAC Date :01/09/26 to 31/10/27"
        orders.cell(7, 11, "Job Number")
        for row_number, job in enumerate(
            ("AA111111MJ", "AA222222MJ", "AA333333MJ", "AA444444MJ", "AA555555MJ", "AA666666MJ", "AA737373MJ", "AA979797MJ"),
            8,
        ):
            orders.cell(row_number, 11, job)
        bds.save(cls.bds_path)

        cls.messages = []
        cls.output_path = Path(processor.process_reconciliation(
            cls.sales_path,
            cls.fotp_path,
            cls.shipment_path,
            cls.bds_path,
            folder,
            cls.messages.append,
        ))
        cls.output = load_workbook(cls.output_path, data_only=False)
        cls.addClassCleanup(cls.output.close)
        cls.october_rows = {
            str(row[0].value): row[0].row
            for row in cls.output["Recon Oct'26"].iter_rows(min_row=4, max_col=1)
            if row[0].value is not None
        }
        cls.november_rows = {
            str(row[0].value): row[0].row
            for row in cls.output["Recon Nov'26"].iter_rows(min_row=4, max_col=1)
            if row[0].value is not None
        }

    def test_new_old_orientation_amounts_and_variance_update(self):
        october = self.output["Recon Oct'26"]
        row = self.october_rows["AA888888MJ"]
        self.assertEqual([october.cell(row, column).value for column in (2, 3, 4, 5)], [10, 120, 10, 100])
        self.assertEqual(october.cell(row, 6).value, f"=B{row}-D{row}")
        self.assertEqual(october.cell(row, 7).value, f"=C{row}-E{row}")
        variance = self.output["45b.Variance FY27"]
        self.assertEqual((variance["C35"].value, variance["D35"].value), (625, 670))
        self.assertEqual((variance["C37"].value, variance["D37"].value), (300, 345))
        self.assertEqual(variance["H28"].value, 20)
        self.assertEqual(variance["C29"].value, 1234)
        self.assertEqual(variance["J28"].value, 777)
        self.assertEqual(self.output["Archive"]["A1"].value, "leave this sheet alone")
        self.assertEqual(self.output["Oct'26"]["A2"].value, "AA888888MJ")
        self.assertNotEqual(self.output_path, self.sales_path)
        source = load_workbook(self.sales_path, read_only=True)
        source_sheetnames = source.sheetnames
        source.close()
        self.assertNotIn("Recon Oct'26", source_sheetnames)
        self.assertNotIn("Recon Nov'26", source_sheetnames)

    def test_remark_precedence_and_cross_month_movements(self):
        october = self.output["Recon Oct'26"]
        expected = {
            "AA888888MJ": "Price Discrepancy",
            "AJ005SSSLS": "Salesman Sample",
            "BA125MJJA26A3": "Demand Pull",
            "AJ089000MP": "Forecast PO",
            "AJ005010MJ": "Forecast PO New",
            "AA111111MJ": "Early ship fr Nov'26 to Oct'26",
            "AA222222MJ": "Delay ship fr Oct'26 to Nov'26",
            "AA333333MJ": "shortship",
            "AA444444MJ": "TBA",
        }
        for job, remark in expected.items():
            with self.subTest(job=job):
                self.assertEqual(october.cell(self.october_rows[job], 8).value, remark)

    def test_forecast_po_new_uses_previous_bds_membership_for_any_variance(self):
        october = self.output["Recon Oct'26"]
        for job, new_qty, old_qty in (("AC123456MJ", 7, 5), ("AC234567MJ", None, 8)):
            with self.subTest(job=job):
                row = self.october_rows[job]
                self.assertEqual(october.cell(row, 2).value, new_qty)
                self.assertEqual(october.cell(row, 4).value, old_qty)
                self.assertEqual(october.cell(row, 8).value, "Forecast PO New")

    def test_fotp_aggregation_conflicting_codes_and_dynamic_comments(self):
        november = self.output["Recon Nov'26"]
        for job, code in (("AA222222MJ", 21), ("AA555555MJ", None), ("AA666666MJ", "#"), ("AA737373MJ", 3), ("AA979797MJ", 31)):
            with self.subTest(job=job):
                row = self.november_rows[job]
                self.assertEqual(november.cell(row, 13).value, code)
                formula = november.cell(row, 14).value
                self.assertIsInstance(formula, str)
                self.assertIn("'Reason code'", formula)
                self.assertIn(f"M{row}", formula)

    def test_delay_variance_buckets_use_gac_codes_not_reason_comments(self):
        variance = self.output["45b.Variance FY27"]
        self.assertEqual(variance["H50"].value, -20)
        self.assertEqual(variance["I50"].value, 20)
        self.assertEqual(variance["H55"].value, -30)
        self.assertEqual(variance["I55"].value, 30)


if __name__ == "__main__":
    unittest.main()
