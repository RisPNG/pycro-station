import hashlib
import importlib.util
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName


spec = importlib.util.spec_from_file_location("sales_forecast_xlsx_package", Path(__file__).with_name("xlsx_package.py"))
package = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = package
spec.loader.exec_module(package)


class ReconciliationPackageTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "before.xlsx"
        self.output = Path(self.directory.name) / "after.xlsx"
        workbook = Workbook()
        variance = workbook.active
        variance.title = "45b.Variance FY27"
        variance["C5"] = "=2+3"
        variance["H28"] = 0
        reason = workbook.create_sheet("Reason code")
        reason["A3"] = 21
        reason["C3"] = "Delay due to development"
        october = workbook.create_sheet("Oct'26")
        october["A1"] = "Job"
        october["B1"] = "Qty"
        october["A2"] = "AJ005001MJ"
        october["B2"] = 29
        october.print_area = "A1:B2"
        chart = BarChart()
        chart.add_data(Reference(october, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        october.add_chart(chart, "D1")
        workbook.create_sheet("Nov'26")
        workbook.save(self.source)
        with zipfile.ZipFile(self.source, "a") as archive:
            archive.writestr("xl/custom/opaque.bin", b"preserve this nonworksheet part")

    def test_adds_recon_sheets_and_preserves_unrelated_parts(self):
        initial_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        package.write_reconciliation_workbook(
            self.source,
            self.output,
            {
                "Recon Oct'26": [
                    {"job": "AJ005001MJ", "new_qty": 29, "new_amount": 400.25, "old_qty": 25, "old_amount": 350.25, "remark": "Delay ship", "gac": 21},
                    {"job": "AJ005002MJ", "new_qty": None, "new_amount": None, "old_qty": 3, "old_amount": 12, "remark": "Missing", "gac": None},
                ],
                "Recon Nov'26": [],
            },
            {"H28": 12.34, "B2": 100},
            ("300926", "230926"),
        )
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), initial_hash)
        with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(self.output) as after:
            self.assertIsNone(after.testzip())
            changed = {
                "[Content_Types].xml",
                "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels",
                "xl/styles.xml",
                "xl/worksheets/sheet1.xml",
            }
            for name in before.namelist():
                if name not in changed:
                    self.assertEqual(after.read(name), before.read(name), name)
            self.assertEqual(after.read("xl/custom/opaque.bin"), b"preserve this nonworksheet part")
        formulas = load_workbook(self.output, data_only=False)
        self.assertEqual(formulas.sheetnames, ["45b.Variance FY27", "Reason code", "Recon Oct'26", "Oct'26", "Recon Nov'26", "Nov'26"])
        self.assertEqual(formulas["45b.Variance FY27"]["C5"].value, "=2+3")
        self.assertEqual(formulas["45b.Variance FY27"]["H28"].value, 12.34)
        self.assertEqual(formulas["45b.Variance FY27"]["B2"].value, 100)
        self.assertEqual(len(formulas["Oct'26"]._charts), 1)
        self.assertEqual(formulas["Oct'26"].print_area, "'Oct''26'!$A$1:$B$2")
        recon = formulas["Recon Oct'26"]
        self.assertEqual(recon["B1"].value, "300926")
        self.assertEqual(recon["D1"].value, "230926")
        self.assertEqual(recon["A3"].value, "Job")
        self.assertEqual(recon["H3"].value, "Remark")
        self.assertEqual(recon["L3"].value, "Remark - SIV Comment")
        self.assertEqual(recon["M3"].value, "GAC CODE")
        self.assertEqual(recon["N3"].value, "SIG Finance Comment")
        self.assertEqual(recon.auto_filter.ref, "A3:N5")
        self.assertEqual(recon["F4"].value, "=B4-D4")
        self.assertEqual(recon["G4"].value, "=C4-E4")
        self.assertIn("VLOOKUP", recon["N4"].value)
        self.assertIn("'Reason code'!$A:$C", recon["N4"].value)
        self.assertEqual(recon["M4"].value, 21)
        self.assertIsNone(recon["M5"].value)
        formulas.close()
        values = load_workbook(self.output, data_only=True, read_only=True)
        self.assertEqual(values["Recon Oct'26"]["F4"].value, 4)
        self.assertEqual(values["Recon Oct'26"]["G4"].value, 50)
        self.assertEqual(values["Recon Oct'26"]["F5"].value, -3)
        self.assertEqual(values["Recon Oct'26"]["B3"].value, 29)
        values.close()
        with zipfile.ZipFile(self.output) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            defined_names = workbook_xml.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedNames")
            print_area = next(node for node in defined_names if node.attrib.get("name") == "_xlnm.Print_Area")
            self.assertEqual(print_area.attrib["localSheetId"], "3")

    def test_replaces_existing_recon_with_optional_ok_suffix(self):
        workbook = load_workbook(self.source)
        old = workbook.create_sheet("Recon Oct'26 ok", 0)
        old["A4"] = "stale job"
        old.print_area = "A3:O3600"
        old.defined_names.add(DefinedName(name="_xlnm._FilterDatabase", attr_text="'Recon Oct''26 ok'!$A$3:$O$3600", localSheetId=0, hidden=True))
        workbook["Nov'26"]["A1"] = "='Recon Oct''26 ok'!B4"
        workbook.save(self.source)
        package.write_reconciliation_workbook(
            self.source,
            self.output,
            {"Recon Oct'26": [{"job": "AJ005001MJ", "new_qty": 29, "new_amount": 400, "old_qty": 29, "old_amount": 400, "remark": "", "gac": "#"}]},
            {},
            ("New", "Old"),
        )
        workbook = load_workbook(self.output, read_only=True)
        self.assertEqual(workbook.sheetnames.count("Recon Oct'26 ok"), 1)
        self.assertNotIn("Recon Oct'26", workbook.sheetnames)
        self.assertLess(workbook.sheetnames.index("Recon Oct'26 ok"), workbook.sheetnames.index("Oct'26"))
        self.assertEqual(workbook["Recon Oct'26 ok"]["A4"].value, "AJ005001MJ")
        self.assertEqual(workbook["Recon Oct'26 ok"]["M4"].value, "#")
        self.assertEqual(workbook["Nov'26"]["A1"].value, "='Recon Oct''26 ok'!B4")
        workbook.close()
        with zipfile.ZipFile(self.output) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            defined_names = workbook_xml.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedNames")
            filters = [node for node in defined_names if node.attrib.get("name") == "_xlnm._FilterDatabase"]
            self.assertEqual(len(filters), 1)
            self.assertEqual(filters[0].text, "'Recon Oct''26 ok'!$A$3:$N$4")
            print_areas = [node for node in defined_names if node.attrib.get("name") == "_xlnm.Print_Area" and node.text.startswith("'Recon Oct''26 ok'!")]
            self.assertEqual(len(print_areas), 1)
            self.assertEqual(print_areas[0].text, "'Recon Oct''26 ok'!$A$3:$N$4")

    def test_does_not_overwrite_source_or_existing_output(self):
        with self.assertRaisesRegex(ValueError, "different"):
            package.write_reconciliation_workbook(self.source, self.source, {}, {}, ("New", "Old"))
        self.output.write_bytes(b"existing")
        with self.assertRaises(FileExistsError):
            package.write_reconciliation_workbook(self.source, self.output, {}, {}, ("New", "Old"))
        self.assertEqual(self.output.read_bytes(), b"existing")

    def test_detects_the_unique_variance_sheet_for_another_fiscal_year(self):
        workbook = load_workbook(self.source)
        workbook["45b.Variance FY27"].title = "45b.Variance FY28"
        workbook.save(self.source)
        package.write_reconciliation_workbook(
            self.source,
            self.output,
            {"Recon Oct'26": []},
            {"H28": 42.5},
            ("New", "Old"),
        )
        output = load_workbook(self.output, read_only=True, data_only=True)
        self.assertEqual(output["45b.Variance FY28"]["H28"].value, 42.5)
        output.close()

    def test_discards_stale_calculation_chain(self):
        with zipfile.ZipFile(self.source) as archive:
            parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
        relationships = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
        ET.SubElement(relationships, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
            "Id": "rId99",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
            "Target": "calcChain.xml",
        })
        parts["xl/_rels/workbook.xml.rels"] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
        content = ET.fromstring(parts["[Content_Types].xml"])
        ET.SubElement(content, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
            "PartName": "/xl/calcChain.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
        })
        parts["[Content_Types].xml"] = ET.tostring(content, encoding="utf-8", xml_declaration=True)
        parts["xl/calcChain.xml"] = b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="N999" i="4"/></calcChain>'
        with zipfile.ZipFile(self.source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

        package.write_reconciliation_workbook(
            self.source,
            self.output,
            {"Recon Oct'26": []},
            {},
            ("New", "Old"),
        )
        with zipfile.ZipFile(self.output) as archive:
            self.assertNotIn("xl/calcChain.xml", archive.namelist())
            self.assertNotIn(b"calcChain", archive.read("xl/_rels/workbook.xml.rels"))
            self.assertNotIn(b"calcChain", archive.read("[Content_Types].xml"))
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            calculation = workbook.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr")
            self.assertEqual(calculation.attrib["fullCalcOnLoad"], "1")

    def test_updates_extended_worksheet_titles_without_losing_named_ranges(self):
        with zipfile.ZipFile(self.source) as archive:
            parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
        parts["docProps/app.xml"] = (
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            '<HeadingPairs><vt:vector size="4" baseType="variant">'
            '<vt:variant><vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>4</vt:i4></vt:variant>'
            '<vt:variant><vt:lpstr>Named Ranges</vt:lpstr></vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant>'
            '</vt:vector></HeadingPairs><TitlesOfParts><vt:vector size="5" baseType="lpstr">'
            '<vt:lpstr>45b.Variance FY27</vt:lpstr><vt:lpstr>Reason code</vt:lpstr>'
            "<vt:lpstr>Oct'26</vt:lpstr><vt:lpstr>Nov'26</vt:lpstr>"
            '<vt:lpstr>Existing Named Range</vt:lpstr></vt:vector></TitlesOfParts></Properties>'
        ).encode()
        with zipfile.ZipFile(self.source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

        package.write_reconciliation_workbook(
            self.source,
            self.output,
            {"Recon Oct'26": [], "Recon Nov'26": []},
            {},
            ("New", "Old"),
        )
        with zipfile.ZipFile(self.output) as archive:
            app = ET.fromstring(archive.read("docProps/app.xml"))
            vt = "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}"
            extended = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
            heading = app.find(f"{extended}HeadingPairs/{vt}vector")
            self.assertEqual(heading[1][0].text, "6")
            titles = app.find(f"{extended}TitlesOfParts/{vt}vector")
            self.assertEqual(titles.attrib["size"], "7")
            self.assertEqual([node.text for node in titles], [
                "45b.Variance FY27", "Reason code", "Recon Oct'26", "Oct'26", "Recon Nov'26", "Nov'26", "Existing Named Range",
            ])


if __name__ == "__main__":
    unittest.main()
